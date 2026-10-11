"""
rapport-auto — Génère un rapport Excel mis en forme à partir d'un export CSV
de ventes, puis l'envoie (optionnellement) par email.

Usage :
    python rapport.py data/ventes_exemple.csv
    python rapport.py data/ventes_exemple.csv -o sortie/rapport.xlsx
    python rapport.py data/ventes_exemple.csv --envoyer
    python rapport.py data/ventes_exemple.csv --envoyer --dry-run
"""

from __future__ import annotations

import argparse
import logging
import os
import smtplib
import sys
from dataclasses import dataclass
from datetime import date
from email.message import EmailMessage
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

log = logging.getLogger("rapport")

COLONNES_REQUISES = ["date", "client", "produit", "quantite", "prix_unitaire"]

# --- Styles -------------------------------------------------------------------
BLEU = "1F4E78"
BLEU_CLAIR = "DDEBF7"
GRIS = "F2F2F2"
FONT_TITRE = Font(name="Calibri", size=16, bold=True, color=BLEU)
FONT_ENTETE = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FILL_ENTETE = PatternFill("solid", fgColor=BLEU)
FILL_ALTERNE = PatternFill("solid", fgColor=GRIS)
FILL_TOTAL = PatternFill("solid", fgColor=BLEU_CLAIR)
BORDURE = Border(bottom=Side(style="thin", color="BFBFBF"))
FORMAT_EUROS = '#,##0.00 "€"'
FORMAT_ENTIER = "#,##0"


# --- Chargement & validation --------------------------------------------------
class DonneesInvalides(ValueError):
    """Le fichier d'entrée ne respecte pas le format attendu."""


def charger_ventes(chemin: Path) -> pd.DataFrame:
    """Lit le CSV, valide les colonnes et nettoie les types.

    Accepte les séparateurs ',' et ';' (exports Excel français) et les
    décimales avec virgule.
    """
    if not chemin.exists():
        raise FileNotFoundError(f"Fichier introuvable : {chemin}")

    df = pd.read_csv(chemin, sep=None, engine="python", dtype=str)
    df.columns = [c.strip().lower() for c in df.columns]

    manquantes = [c for c in COLONNES_REQUISES if c not in df.columns]
    if manquantes:
        raise DonneesInvalides(
            f"Colonnes manquantes : {', '.join(manquantes)}. "
            f"Colonnes attendues : {', '.join(COLONNES_REQUISES)}"
        )

    df = df[COLONNES_REQUISES].copy()
    for col in ("client", "produit"):
        df[col] = df[col].str.strip()

    df["date"] = pd.to_datetime(df["date"], errors="coerce", dayfirst=True, format="mixed")
    for col in ("quantite", "prix_unitaire"):
        df[col] = pd.to_numeric(
            df[col].str.replace(",", ".", regex=False).str.replace(" ", "", regex=False),
            errors="coerce",
        )

    invalides = df[df.isna().any(axis=1)]
    if not invalides.empty:
        log.warning("%d ligne(s) ignorée(s) car incomplètes ou invalides.", len(invalides))
        df = df.dropna()

    if df.empty:
        raise DonneesInvalides("Aucune ligne exploitable dans le fichier.")

    df["montant"] = df["quantite"] * df["prix_unitaire"]
    df["mois"] = df["date"].dt.to_period("M").astype(str)
    return df.sort_values("date").reset_index(drop=True)


# --- Calculs ------------------------------------------------------------------
@dataclass
class Synthese:
    ca_total: float
    nb_ventes: int
    nb_clients: int
    panier_moyen: float
    periode_debut: date
    periode_fin: date
    par_mois: pd.DataFrame
    par_produit: pd.DataFrame
    top_clients: pd.DataFrame


def calculer_synthese(df: pd.DataFrame, nb_top: int = 5) -> Synthese:
    par_mois = (
        df.groupby("mois", as_index=False)
        .agg(ventes=("montant", "size"), ca=("montant", "sum"))
        .sort_values("mois")
    )
    par_produit = (
        df.groupby("produit", as_index=False)
        .agg(quantite=("quantite", "sum"), ca=("montant", "sum"))
        .sort_values("ca", ascending=False)
    )
    par_produit["part"] = par_produit["ca"] / par_produit["ca"].sum()

    top_clients = (
        df.groupby("client", as_index=False)
        .agg(commandes=("montant", "size"), ca=("montant", "sum"))
        .sort_values("ca", ascending=False)
        .head(nb_top)
    )

    ca_total = float(df["montant"].sum())
    return Synthese(
        ca_total=ca_total,
        nb_ventes=len(df),
        nb_clients=df["client"].nunique(),
        panier_moyen=ca_total / len(df),
        periode_debut=df["date"].min().date(),
        periode_fin=df["date"].max().date(),
        par_mois=par_mois.reset_index(drop=True),
        par_produit=par_produit.reset_index(drop=True),
        top_clients=top_clients.reset_index(drop=True),
    )


# --- Écriture Excel -----------------------------------------------------------
def _ecrire_tableau(
    ws: Worksheet,
    ligne: int,
    col: int,
    entetes: list[str],
    lignes: list[list],
    formats: list[str | None],
    total: list | None = None,
) -> int:
    """Écrit un tableau stylé et renvoie la ligne qui suit."""
    for j, titre in enumerate(entetes):
        c = ws.cell(row=ligne, column=col + j, value=titre)
        c.font, c.fill = FONT_ENTETE, FILL_ENTETE
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, valeurs in enumerate(lignes, start=1):
        for j, v in enumerate(valeurs):
            c = ws.cell(row=ligne + i, column=col + j, value=v)
            if formats[j]:
                c.number_format = formats[j]
            c.border = BORDURE
            if i % 2 == 0:
                c.fill = FILL_ALTERNE

    suivante = ligne + len(lignes) + 1
    if total:
        for j, v in enumerate(total):
            c = ws.cell(row=suivante, column=col + j, value=v)
            c.font, c.fill = Font(bold=True), FILL_TOTAL
            if formats[j] and v is not None and not isinstance(v, str):
                c.number_format = formats[j]
        suivante += 1
    return suivante


def _ajuster_largeurs(ws: Worksheet, minimum: int = 10, maximum: int = 40) -> None:
    largeurs: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            taille = len(f"{c.value:,.2f}") + 3 if isinstance(c.value, float) else len(str(c.value))
            largeurs[c.column] = max(largeurs.get(c.column, 0), taille)
    for col, l in largeurs.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minimum, min(maximum, l + 2))


def _feuille_synthese(wb: Workbook, s: Synthese, titre: str) -> None:
    ws = wb.active
    ws.title = "Synthèse"
    ws.sheet_view.showGridLines = False

    ws["A1"] = titre
    ws["A1"].font = FONT_TITRE
    ws["A2"] = (
        f"Période du {s.periode_debut:%d/%m/%Y} au {s.periode_fin:%d/%m/%Y} "
        f"— généré le {date.today():%d/%m/%Y}"
    )
    ws["A2"].font = Font(italic=True, color="7F7F7F")

    # Indicateurs clés
    indicateurs = [
        ("Chiffre d'affaires", s.ca_total, FORMAT_EUROS),
        ("Nombre de ventes", s.nb_ventes, FORMAT_ENTIER),
        ("Clients distincts", s.nb_clients, FORMAT_ENTIER),
        ("Panier moyen", s.panier_moyen, FORMAT_EUROS),
    ]
    for j, (libelle, valeur, fmt) in enumerate(indicateurs):
        col = 1 + j
        lib = ws.cell(row=4, column=col, value=libelle)
        lib.font = Font(size=9, color="7F7F7F")
        val = ws.cell(row=5, column=col, value=valeur)
        val.font = Font(size=14, bold=True, color=BLEU)
        val.number_format = fmt

    # CA par mois
    ws.cell(row=7, column=1, value="Évolution mensuelle").font = Font(size=12, bold=True)
    debut_mois = 8
    fin = _ecrire_tableau(
        ws, debut_mois, 1,
        ["Mois", "Ventes", "CA"],
        s.par_mois[["mois", "ventes", "ca"]].values.tolist(),
        [None, FORMAT_ENTIER, FORMAT_EUROS],
        total=["Total", s.nb_ventes, s.ca_total],
    )

    graph = LineChart()
    graph.title = "Chiffre d'affaires mensuel"
    graph.y_axis.title = "€"
    graph.height, graph.width = 7.5, 16
    graph.legend = None
    data = Reference(ws, min_col=3, min_row=debut_mois, max_row=debut_mois + len(s.par_mois))
    cats = Reference(ws, min_col=1, min_row=debut_mois + 1, max_row=debut_mois + len(s.par_mois))
    graph.add_data(data, titles_from_data=True)
    graph.set_categories(cats)
    graph.series[0].smooth = False
    ws.add_chart(graph, "F7")

    # Produits
    ligne = max(fin, 24) + 1
    ws.cell(row=ligne, column=1, value="Ventes par produit").font = Font(size=12, bold=True)
    debut_prod = ligne + 1
    fin = _ecrire_tableau(
        ws, debut_prod, 1,
        ["Produit", "Quantité", "CA", "Part du CA"],
        s.par_produit[["produit", "quantite", "ca", "part"]].values.tolist(),
        [None, FORMAT_ENTIER, FORMAT_EUROS, "0.0%"],
    )

    barres = BarChart()
    barres.type = "bar"
    barres.title = "CA par produit"
    barres.height, barres.width = 7.5, 16
    barres.legend = None
    data = Reference(ws, min_col=3, min_row=debut_prod, max_row=debut_prod + len(s.par_produit))
    cats = Reference(ws, min_col=1, min_row=debut_prod + 1, max_row=debut_prod + len(s.par_produit))
    barres.add_data(data, titles_from_data=True)
    barres.set_categories(cats)
    ws.add_chart(barres, f"F{ligne}")

    # Top clients
    ligne = max(fin, ligne + 16) + 1
    ws.cell(row=ligne, column=1, value=f"Top {len(s.top_clients)} clients").font = Font(size=12, bold=True)
    _ecrire_tableau(
        ws, ligne + 1, 1,
        ["Client", "Commandes", "CA"],
        s.top_clients[["client", "commandes", "ca"]].values.tolist(),
        [None, FORMAT_ENTIER, FORMAT_EUROS],
    )

    _ajuster_largeurs(ws, minimum=14)
    _mise_en_page(ws, paysage=True)


def _mise_en_page(ws: Worksheet, paysage: bool = False) -> None:
    """Impression propre : tient en largeur sur une page A4."""
    ws.page_setup.orientation = "landscape" if paysage else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def _feuille_detail(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Détail des ventes")
    lignes = [
        [r.date.date(), r.client, r.produit, r.quantite, r.prix_unitaire, r.montant]
        for r in df.itertuples()
    ]
    _ecrire_tableau(
        ws, 1, 1,
        ["Date", "Client", "Produit", "Quantité", "Prix unitaire", "Montant"],
        lignes,
        ["DD/MM/YYYY", None, None, FORMAT_ENTIER, FORMAT_EUROS, FORMAT_EUROS],
    )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(lignes) + 1}"
    ws.print_title_rows = "1:1"
    _ajuster_largeurs(ws)
    _mise_en_page(ws)


def generer_excel(df: pd.DataFrame, s: Synthese, sortie: Path, titre: str) -> Path:
    sortie.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _feuille_synthese(wb, s, titre)
    _feuille_detail(wb, df)
    wb.save(sortie)
    return sortie


# --- Email --------------------------------------------------------------------
@dataclass
class ConfigSMTP:
    hote: str
    port: int
    utilisateur: str
    mot_de_passe: str
    expediteur: str
    destinataires: list[str]

    @classmethod
    def depuis_env(cls) -> "ConfigSMTP":
        requis = ["SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "MAIL_TO"]
        manquants = [v for v in requis if not os.getenv(v)]
        if manquants:
            raise RuntimeError(
                "Variables d'environnement manquantes pour l'envoi : "
                + ", ".join(manquants)
                + " (voir .env.example)"
            )
        return cls(
            hote=os.environ["SMTP_HOST"],
            port=int(os.getenv("SMTP_PORT", "587")),
            utilisateur=os.environ["SMTP_USER"],
            mot_de_passe=os.environ["SMTP_PASSWORD"],
            expediteur=os.getenv("MAIL_FROM", os.environ["SMTP_USER"]),
            destinataires=[a.strip() for a in os.environ["MAIL_TO"].split(",") if a.strip()],
        )


def euros(valeur: float) -> str:
    """Format monétaire français : 28 680,00 €"""
    return f"{valeur:,.2f} €".replace(",", "\u202f").replace(".", ",")


def construire_email(s: Synthese, piece_jointe: Path, cfg: ConfigSMTP, titre: str) -> EmailMessage:
    msg = EmailMessage()
    msg["Subject"] = f"{titre} — {s.periode_debut:%d/%m/%Y} au {s.periode_fin:%d/%m/%Y}"
    msg["From"] = cfg.expediteur
    msg["To"] = ", ".join(cfg.destinataires)

    top = s.par_produit.iloc[0]
    msg.set_content(
        "Bonjour,\n\n"
        f"Voici le rapport des ventes du {s.periode_debut:%d/%m/%Y} au {s.periode_fin:%d/%m/%Y}.\n\n"
        f"  • Chiffre d'affaires : {euros(s.ca_total)}\n"
        f"  • Nombre de ventes   : {s.nb_ventes}\n"
        f"  • Panier moyen       : {euros(s.panier_moyen)}\n"
        f"  • Produit n°1        : {top['produit']} ({top['part']:.0%} du CA)\n\n"
        "Le détail complet est dans le fichier Excel joint.\n\n"
        "Ce message a été généré automatiquement.\n"
    )
    msg.add_attachment(
        piece_jointe.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=piece_jointe.name,
    )
    return msg


def envoyer_email(msg: EmailMessage, cfg: ConfigSMTP) -> None:
    if cfg.port == 465:
        with smtplib.SMTP_SSL(cfg.hote, cfg.port, timeout=30) as smtp:
            smtp.login(cfg.utilisateur, cfg.mot_de_passe)
            smtp.send_message(msg)
    else:
        with smtplib.SMTP(cfg.hote, cfg.port, timeout=30) as smtp:
            smtp.starttls()
            smtp.login(cfg.utilisateur, cfg.mot_de_passe)
            smtp.send_message(msg)


# --- CLI ----------------------------------------------------------------------
def _charger_dotenv(chemin: Path = Path(".env")) -> None:
    """Charge un fichier .env simple (CLE=valeur) sans dépendance externe."""
    if not chemin.exists():
        return
    for ligne in chemin.read_text(encoding="utf-8").splitlines():
        ligne = ligne.strip()
        if not ligne or ligne.startswith("#") or "=" not in ligne:
            continue
        cle, valeur = ligne.split("=", 1)
        os.environ.setdefault(cle.strip(), valeur.strip().strip('"').strip("'"))


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Génère un rapport Excel de ventes à partir d'un CSV.")
    p.add_argument("csv", type=Path, help="Fichier CSV d'entrée")
    p.add_argument("-o", "--sortie", type=Path, help="Fichier Excel de sortie (défaut : rapports/rapport_AAAA-MM-JJ.xlsx)")
    p.add_argument("-t", "--titre", default="Rapport des ventes", help="Titre du rapport")
    p.add_argument("--envoyer", action="store_true", help="Envoyer le rapport par email (config dans .env)")
    p.add_argument("--dry-run", action="store_true", help="Avec --envoyer : prépare l'email sans l'envoyer")
    p.add_argument("-v", "--verbose", action="store_true")
    args = p.parse_args(argv)

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO, format="%(levelname)s  %(message)s")
    _charger_dotenv()

    try:
        df = charger_ventes(args.csv)
        s = calculer_synthese(df)
        sortie = args.sortie or Path("rapports") / f"rapport_{date.today():%Y-%m-%d}.xlsx"
        generer_excel(df, s, sortie, args.titre)
        log.info("Rapport généré : %s (%d ventes, CA %.2f €)", sortie, s.nb_ventes, s.ca_total)

        if args.envoyer:
            cfg = ConfigSMTP.depuis_env()
            msg = construire_email(s, sortie, cfg, args.titre)
            if args.dry_run:
                log.info("[dry-run] Email prêt pour %s — non envoyé.", msg["To"])
                print("\n" + msg.get_body(preferencelist=("plain",)).get_content())
            else:
                envoyer_email(msg, cfg)
                log.info("Email envoyé à %s", msg["To"])
    except (FileNotFoundError, DonneesInvalides, RuntimeError) as e:
        log.error("%s", e)
        return 1
    except smtplib.SMTPException as e:
        log.error("Échec de l'envoi de l'email : %s", e)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
