from email.message import EmailMessage
from pathlib import Path

import pytest
from openpyxl import load_workbook

import sales_report

VALID_CSV = """date,customer,product,quantity,unit_price
2026-01-15,Customer A,Product X,2,10.50
2026-01-20,Customer B,Product Y,1,100
2026-02-03,Customer A,Product X,3,10.50
"""


@pytest.fixture
def valid_csv(tmp_path: Path) -> Path:
    f = tmp_path / "sales.csv"
    f.write_text(VALID_CSV, encoding="utf-8")
    return f


def test_load_and_amounts(valid_csv):
    df = sales_report.load_sales(valid_csv)
    assert len(df) == 3
    assert df["amount"].sum() == pytest.approx(2 * 10.5 + 100 + 3 * 10.5)
    assert list(df["month"].unique()) == ["2026-01", "2026-02"]


def test_european_format(tmp_path):
    f = tmp_path / "s.csv"
    f.write_text("date;customer;product;quantity;unit_price\n01/03/2026;A;X;1;9,99\n", encoding="utf-8")
    df = sales_report.load_sales(f)
    assert df["amount"].iloc[0] == pytest.approx(9.99)
    assert df["date"].iloc[0].month == 3  # day-first: 1 March, not 3 January


@pytest.mark.parametrize("text, expected", [
    ("2026-02-03", (2026, 2, 3)),      # ISO: never swapped
    ("03/02/2026", (2026, 2, 3)),      # European day-first
    ("2026-12-01", (2026, 12, 1)),
    ("31/12/2026", (2026, 12, 31)),
])
def test_parse_dates(text, expected):
    d = sales_report.parse_dates(sales_report.pd.Series([text])).iloc[0]
    assert (d.year, d.month, d.day) == expected


def test_parse_dates_mixed_and_invalid():
    s = sales_report.parse_dates(sales_report.pd.Series(["2026-01-05", "05/01/2026", "nonsense", ""]))
    assert s.iloc[0] == s.iloc[1]
    assert s.iloc[2:].isna().all()


def test_missing_columns(tmp_path):
    f = tmp_path / "s.csv"
    f.write_text("date,customer,quantity\n2026-01-01,A,1\n", encoding="utf-8")
    with pytest.raises(sales_report.InvalidData, match="product"):
        sales_report.load_sales(f)


def test_invalid_rows_skipped(tmp_path):
    f = tmp_path / "s.csv"
    f.write_text(VALID_CSV + "not-a-date,C,Z,1,5\n2026-02-10,D,Z,abc,5\n", encoding="utf-8")
    df = sales_report.load_sales(f)
    assert len(df) == 3


def test_file_not_found(tmp_path):
    with pytest.raises(FileNotFoundError):
        sales_report.load_sales(tmp_path / "missing.csv")


def test_summary(valid_csv):
    s = sales_report.summarize(sales_report.load_sales(valid_csv))
    assert s.sales_count == 3
    assert s.customer_count == 2
    assert s.by_product.iloc[0]["product"] == "Product Y"
    assert s.by_product["share"].sum() == pytest.approx(1.0)
    assert s.top_customers.iloc[0]["customer"] == "Customer B"


def test_build_excel(valid_csv, tmp_path):
    df = sales_report.load_sales(valid_csv)
    s = sales_report.summarize(df)
    output = sales_report.build_excel(df, s, tmp_path / "out" / "r.xlsx", "Test")
    wb = load_workbook(output)
    assert wb.sheetnames == ["Summary", "Sales detail"]
    assert wb["Summary"]["A1"].value == "Test"
    assert wb["Sales detail"].max_row == 4
    assert len(wb["Summary"]._charts) == 2


def test_email(valid_csv, tmp_path):
    df = sales_report.load_sales(valid_csv)
    s = sales_report.summarize(df)
    xlsx = sales_report.build_excel(df, s, tmp_path / "r.xlsx", "Test")
    cfg = sales_report.SMTPConfig("smtp.test", 587, "u", "p", "me@test.com", ["a@test.com", "b@test.com"])
    msg = sales_report.build_email(s, xlsx, cfg, "Test")
    assert isinstance(msg, EmailMessage)
    assert msg["To"] == "a@test.com, b@test.com"
    attachments = list(msg.iter_attachments())
    assert len(attachments) == 1 and attachments[0].get_filename() == "r.xlsx"


def test_missing_env_config(monkeypatch):
    for v in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "MAIL_TO"):
        monkeypatch.delenv(v, raising=False)
    with pytest.raises(RuntimeError, match="SMTP_HOST"):
        sales_report.SMTPConfig.from_env()


def test_cli_dry_run(valid_csv, tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("SMTP_HOST", "smtp.test")
    monkeypatch.setenv("SMTP_USER", "u@test.com")
    monkeypatch.setenv("SMTP_PASSWORD", "x")
    monkeypatch.setenv("MAIL_TO", "a@test.com")
    code = sales_report.main([str(valid_csv), "-o", "r.xlsx", "--send", "--dry-run"])
    assert code == 0
    assert (tmp_path / "r.xlsx").exists()
    assert "Revenue" in capsys.readouterr().out


def test_money_format():
    assert sales_report.money(28680) == "€28,680.00"
    assert sales_report.money(9.5) == "€9.50"
