from email.message import EmailMessage
from pathlib import Path

import pytest
from openpyxl import load_workbook

import rapport

CSV_VALIDE = """date;client;produit;quantite;prix_unitaire
15/01/2026;Client A;Produit X;2;10,50
20/01/2026;Client B;Produit Y;1;100
03/02/2026;Client A;Produit X;3;10,50
"""


@pytest.fixture
def csv_valide(tmp_path: Path) -> Path:
    f = tmp_path / "ventes.csv"
    f.write_text(CSV_VALIDE, encoding="utf-8")
    return f


def test_chargement_et_montants(csv_valide):
    df = rapport.charger_ventes(csv_valide)
    assert len(df) == 3
    assert df["montant"].sum() == pytest.approx(2 * 10.5 + 100 + 3 * 10.5)
    assert list(df["mois"].unique()) == ["2026-01", "2026-02"]


def test_separateur_virgule(tmp_path):
    f = tmp_path / "v.csv"
    f.write_text("date,client,produit,quantite,prix_unitaire\n2026-03-01,A,X,1,9.99\n", encoding="utf-8")
    df = rapport.charger_ventes(f)
    assert df["montant"].iloc[0] == pytest.approx(9.99)


def test_colonnes_manquantes(tmp_path):
    f = tmp_path / "v.csv"
    f.write_text("date;client;quantite\n01/01/2026;A;1\n", encoding="utf-8")
    with pytest.raises(rapport.DonneesInvalides, match="produit"):
        rapport.charger_ventes(f)


def test_lignes_invalides_ignorees(tmp_path):
    f = tmp_path / "v.csv"
    f.write_text(CSV_VALIDE + "pas-une-date;C;Z;1;5\n10/02/2026;D;Z;abc;5\n", encoding="utf-8")
    df = rapport.charger_ventes(f)
    assert len(df) == 3


def test_fichier_introuvable(tmp_path):
    with pytest.raises(FileNotFoundError):
        rapport.charger_ventes(tmp_path / "absent.csv")


def test_synthese(csv_valide):
    s = rapport.calculer_synthese(rapport.charger_ventes(csv_valide))
    assert s.nb_ventes == 3
    assert s.nb_clients == 2
    assert s.par_produit.iloc[0]["produit"] == "Produit Y"
    assert s.par_produit["part"].sum() == pytest.approx(1.0)
    assert s.top_clients.iloc[0]["client"] == "Client B"


def test_generation_excel(csv_valide, tmp_path):
    df = rapport.charger_ventes(csv_valide)
    s = rapport.calculer_synthese(df)
    sortie = rapport.generer_excel(df, s, tmp_path / "out" / "r.xlsx", "Test")
    wb = load_workbook(sortie)
    assert wb.sheetnames == ["Synthèse", "Détail des ventes"]
    assert wb["Synthèse"]["A1"].value == "Test"
    assert wb["Détail des ventes"].max_row == 4
    assert len(wb["Synthèse"]._charts) == 2


def test_email(csv_valide, tmp_path):
    df = rapport.charger_ventes(csv_valide)
    s = rapport.calculer_synthese(df)
    xlsx = rapport.generer_excel(df, s, tmp_path / "r.xlsx", "Test")
    cfg = rapport.ConfigSMTP("smtp.test", 587, "u", "p", "moi@test.fr", ["a@test.fr", "b@test.fr"])
    msg = rapport.construire_email(s, xlsx, cfg, "Test")
    assert isinstance(msg, EmailMessage)
    assert msg["To"] == "a@test.fr, b@test.fr"
    pj = list(msg.iter_attachments())
    assert len(pj) == 1 and pj[0].get_filename() == "r.xlsx"


def test_config_env_manquante(monkeypatch):
    for v in ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "MAIL_TO"):
        monkeypatch.delenv(v, raising=False)
    with pytest.raises(RuntimeError, match="SMTP_HOST"):
        rapport.ConfigSMTP.depuis_env()


def test_cli_dry_run(csv_valide, tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("SMTP_HOST", "smtp.test")
    monkeypatch.setenv("SMTP_USER", "u@test.fr")
    monkeypatch.setenv("SMTP_PASSWORD", "x")
    monkeypatch.setenv("MAIL_TO", "a@test.fr")
    code = rapport.main([str(csv_valide), "-o", "r.xlsx", "--envoyer", "--dry-run"])
    assert code == 0
    assert (tmp_path / "r.xlsx").exists()
    assert "Chiffre d'affaires" in capsys.readouterr().out


def test_format_euros():
    assert rapport.euros(28680) == "28 680,00 €"
    assert rapport.euros(9.5) == "9,50 €"
