"""
sales-report-mailer — Turn a CSV sales export into a formatted Excel report,
then (optionally) email it.

Usage:
    python sales_report.py data/sample_sales.csv
    python sales_report.py data/sample_sales.csv -o output/report.xlsx
    python sales_report.py data/sample_sales.csv --send
    python sales_report.py data/sample_sales.csv --send --dry-run
"""

from __future__ import annotations

import argparse
import logging
import os
import smtplib
import sys
from dataclasses import dataclass
from datetime import date
from email.message import EmailMessage
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

log = logging.getLogger("sales_report")

REQUIRED_COLUMNS = ["date", "customer", "product", "quantity", "unit_price"]

# --- Styles -------------------------------------------------------------------
BLUE = "1F4E78"
LIGHT_BLUE = "DDEBF7"
GREY = "F2F2F2"
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=BLUE)
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=BLUE)
STRIPE_FILL = PatternFill("solid", fgColor=GREY)
TOTAL_FILL = PatternFill("solid", fgColor=LIGHT_BLUE)
BOTTOM_BORDER = Border(bottom=Side(style="thin", color="BFBFBF"))
MONEY_FORMAT = '"€"#,##0.00'
INT_FORMAT = "#,##0"


# --- Loading & validation -----------------------------------------------------
class InvalidData(ValueError):
    """The input file does not match the expected format."""


def parse_dates(values: pd.Series) -> pd.Series:
    """Parse ISO dates (2026-02-03) as-is and other dates as day-first (03/02/2026).

    Applying day-first parsing to everything would silently swap the day and
    month of ISO dates, so the two formats are handled separately.
    """
    values = values.str.strip()
    iso = values.str.match(r"^\d{4}-\d{1,2}-\d{1,2}")
    parsed = pd.Series(pd.NaT, index=values.index, dtype="datetime64[ns]")
    parsed[iso] = pd.to_datetime(values[iso], errors="coerce", format="ISO8601")
    parsed[~iso] = pd.to_datetime(values[~iso], errors="coerce", dayfirst=True, format="mixed")
    return parsed


def load_sales(path: Path) -> pd.DataFrame:
    """Read the CSV, check its columns and clean up the types.

    Accepts both ',' and ';' separators (European Excel exports) and decimal
    commas. Dates may be ISO (2026-04-01) or day-first (01/04/2026).
    """
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    df = pd.read_csv(path, sep=None, engine="python", dtype=str)
    df.columns = [c.strip().lower() for c in df.columns]

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise InvalidData(
            f"Missing columns: {', '.join(missing)}. "
            f"Expected columns: {', '.join(REQUIRED_COLUMNS)}"
        )

    df = df[REQUIRED_COLUMNS].copy()
    for col in ("customer", "product"):
        df[col] = df[col].str.strip()

    df["date"] = parse_dates(df["date"])
    for col in ("quantity", "unit_price"):
        df[col] = pd.to_numeric(
            df[col].str.replace(",", ".", regex=False).str.replace(" ", "", regex=False),
            errors="coerce",
        )

    invalid = df[df.isna().any(axis=1)]
    if not invalid.empty:
        log.warning("%d row(s) skipped because they are incomplete or invalid.", len(invalid))
        df = df.dropna()

    if df.empty:
        raise InvalidData("No usable rows in the file.")

    df["amount"] = df["quantity"] * df["unit_price"]
    df["month"] = df["date"].dt.to_period("M").astype(str)
    return df.sort_values("date").reset_index(drop=True)


# --- Calculations -------------------------------------------------------------
@dataclass
class Summary:
    revenue: float
    sales_count: int
    customer_count: int
    average_order: float
    period_start: date
    period_end: date
    by_month: pd.DataFrame
    by_product: pd.DataFrame
    top_customers: pd.DataFrame


def summarize(df: pd.DataFrame, top_n: int = 5) -> Summary:
    by_month = (
        df.groupby("month", as_index=False)
        .agg(sales=("amount", "size"), revenue=("amount", "sum"))
        .sort_values("month")
    )
    by_product = (
        df.groupby("product", as_index=False)
        .agg(quantity=("quantity", "sum"), revenue=("amount", "sum"))
        .sort_values("revenue", ascending=False)
    )
    by_product["share"] = by_product["revenue"] / by_product["revenue"].sum()

    top_customers = (
        df.groupby("customer", as_index=False)
        .agg(orders=("amount", "size"), revenue=("amount", "sum"))
        .sort_values("revenue", ascending=False)
        .head(top_n)
    )

    revenue = float(df["amount"].sum())
    return Summary(
        revenue=revenue,
        sales_count=len(df),
        customer_count=df["customer"].nunique(),
        average_order=revenue / len(df),
        period_start=df["date"].min().date(),
        period_end=df["date"].max().date(),
        by_month=by_month.reset_index(drop=True),
        by_product=by_product.reset_index(drop=True),
        top_customers=top_customers.reset_index(drop=True),
    )


# --- Excel output -------------------------------------------------------------
def _write_table(
    ws: Worksheet,
    row: int,
    col: int,
    headers: list[str],
    rows: list[list],
    formats: list[str | None],
    total: list | None = None,
) -> int:
    """Write a styled table and return the next free row."""
    for j, title in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=title)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, values in enumerate(rows, start=1):
        for j, v in enumerate(values):
            c = ws.cell(row=row + i, column=col + j, value=v)
            if formats[j]:
                c.number_format = formats[j]
            c.border = BOTTOM_BORDER
            if i % 2 == 0:
                c.fill = STRIPE_FILL

    next_row = row + len(rows) + 1
    if total:
        for j, v in enumerate(total):
            c = ws.cell(row=next_row, column=col + j, value=v)
            c.font, c.fill = Font(bold=True), TOTAL_FILL
            if formats[j] and v is not None and not isinstance(v, str):
                c.number_format = formats[j]
        next_row += 1
    return next_row


def _fit_columns(ws: Worksheet, minimum: int = 10, maximum: int = 40) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            size = len(f"{c.value:,.2f}") + 3 if isinstance(c.value, float) else len(str(c.value))
            widths[c.column] = max(widths.get(c.column, 0), size)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minimum, min(maximum, w + 2))


def _page_setup(ws: Worksheet, landscape: bool = False) -> None:
    """Clean printing: fits one A4 page wide."""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def _summary_sheet(wb: Workbook, s: Summary, title: str) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False

    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"Period: {s.period_start:%Y-%m-%d} to {s.period_end:%Y-%m-%d} "
        f"— generated on {date.today():%Y-%m-%d}"
    )
    ws["A2"].font = Font(italic=True, color="7F7F7F")

    # Key figures
    kpis = [
        ("Revenue", s.revenue, MONEY_FORMAT),
        ("Number of sales", s.sales_count, INT_FORMAT),
        ("Unique customers", s.customer_count, INT_FORMAT),
        ("Average order", s.average_order, MONEY_FORMAT),
    ]
    for j, (label, value, fmt) in enumerate(kpis):
        col = 1 + j
        lab = ws.cell(row=4, column=col, value=label)
        lab.font = Font(size=9, color="7F7F7F")
        val = ws.cell(row=5, column=col, value=value)
        val.font = Font(size=14, bold=True, color=BLUE)
        val.number_format = fmt

    # Revenue by month
    ws.cell(row=7, column=1, value="Monthly trend").font = Font(size=12, bold=True)
    month_start = 8
    end = _write_table(
        ws, month_start, 1,
        ["Month", "Sales", "Revenue"],
        s.by_month[["month", "sales", "revenue"]].values.tolist(),
        [None, INT_FORMAT, MONEY_FORMAT],
        total=["Total", s.sales_count, s.revenue],
    )

    line = LineChart()
    line.title = "Monthly revenue"
    line.y_axis.title = "€"
    line.height, line.width = 7.5, 16
    line.legend = None
    data = Reference(ws, min_col=3, min_row=month_start, max_row=month_start + len(s.by_month))
    cats = Reference(ws, min_col=1, min_row=month_start + 1, max_row=month_start + len(s.by_month))
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    line.series[0].smooth = False
    ws.add_chart(line, "F7")

    # Products
    row = max(end, 24) + 1
    ws.cell(row=row, column=1, value="Sales by product").font = Font(size=12, bold=True)
    product_start = row + 1
    end = _write_table(
        ws, product_start, 1,
        ["Product", "Quantity", "Revenue", "Share of revenue"],
        s.by_product[["product", "quantity", "revenue", "share"]].values.tolist(),
        [None, INT_FORMAT, MONEY_FORMAT, "0.0%"],
    )

    bars = BarChart()
    bars.type = "bar"
    bars.title = "Revenue by product"
    bars.height, bars.width = 7.5, 16
    bars.legend = None
    data = Reference(ws, min_col=3, min_row=product_start, max_row=product_start + len(s.by_product))
    cats = Reference(ws, min_col=1, min_row=product_start + 1, max_row=product_start + len(s.by_product))
    bars.add_data(data, titles_from_data=True)
    bars.set_categories(cats)
    ws.add_chart(bars, f"F{row}")

    # Top customers
    row = max(end, row + 16) + 1
    ws.cell(row=row, column=1, value=f"Top {len(s.top_customers)} customers").font = Font(size=12, bold=True)
    _write_table(
        ws, row + 1, 1,
        ["Customer", "Orders", "Revenue"],
        s.top_customers[["customer", "orders", "revenue"]].values.tolist(),
        [None, INT_FORMAT, MONEY_FORMAT],
    )

    _fit_columns(ws, minimum=14)
    _page_setup(ws, landscape=True)


def _detail_sheet(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Sales detail")
    rows = [
        [r.date.date(), r.customer, r.product, r.quantity, r.unit_price, r.amount]
        for r in df.itertuples()
    ]
    _write_table(
        ws, 1, 1,
        ["Date", "Customer", "Product", "Quantity", "Unit price", "Amount"],
        rows,
        ["YYYY-MM-DD", None, None, INT_FORMAT, MONEY_FORMAT, MONEY_FORMAT],
    )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(rows) + 1}"
    ws.print_title_rows = "1:1"
    _fit_columns(ws)
    _page_setup(ws)


def build_excel(df: pd.DataFrame, s: Summary, output: Path, title: str) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _summary_sheet(wb, s, title)
    _detail_sheet(wb, df)
    wb.save(output)
    return output


# --- Email --------------------------------------------------------------------
@dataclass
class SMTPConfig:
    host: str
    port: int
    user: str
    password: str
    sender: str
    recipients: list[str]

    @classmethod
    def from_env(cls) -> "SMTPConfig":
        required = ["SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "MAIL_TO"]
        missing = [v for v in required if not os.getenv(v)]
        if missing:
            raise RuntimeError(
                "Missing environment variables for sending: "
                + ", ".join(missing)
                + " (see .env.example)"
            )
        return cls(
            host=os.environ["SMTP_HOST"],
            port=int(os.getenv("SMTP_PORT", "587")),
            user=os.environ["SMTP_USER"],
            password=os.environ["SMTP_PASSWORD"],
            sender=os.getenv("MAIL_FROM", os.environ["SMTP_USER"]),
            recipients=[a.strip() for a in os.environ["MAIL_TO"].split(",") if a.strip()],
        )


def money(value: float) -> str:
    """Format an amount in euros: €28,680.00"""
    return f"€{value:,.2f}"


def build_email(s: Summary, attachment: Path, cfg: SMTPConfig, title: str) -> EmailMessage:
    msg = EmailMessage()
    msg["Subject"] = f"{title} — {s.period_start:%Y-%m-%d} to {s.period_end:%Y-%m-%d}"
    msg["From"] = cfg.sender
    msg["To"] = ", ".join(cfg.recipients)

    top = s.by_product.iloc[0]
    msg.set_content(
        "Hello,\n\n"
        f"Here is the sales report from {s.period_start:%Y-%m-%d} to {s.period_end:%Y-%m-%d}.\n\n"
        f"  • Revenue         : {money(s.revenue)}\n"
        f"  • Number of sales : {s.sales_count}\n"
        f"  • Average order   : {money(s.average_order)}\n"
        f"  • Top product     : {top['product']} ({top['share']:.0%} of revenue)\n\n"
        "The full details are in the attached Excel file.\n\n"
        "This message was generated automatically.\n"
    )
    msg.add_attachment(
        attachment.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=attachment.name,
    )
    return msg


def send_email(msg: EmailMessage, cfg: SMTPConfig) -> None:
    if cfg.port == 465:
        with smtplib.SMTP_SSL(cfg.host, cfg.port, timeout=30) as smtp:
            smtp.login(cfg.user, cfg.password)
            smtp.send_message(msg)
    else:
        with smtplib.SMTP(cfg.host, cfg.port, timeout=30) as smtp:
            smtp.starttls()
            smtp.login(cfg.user, cfg.password)
            smtp.send_message(msg)


# --- CLI ----------------------------------------------------------------------
def _load_dotenv(path: Path = Path(".env")) -> None:
    """Load a simple .env file (KEY=value) without any external dependency."""
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Build an Excel sales report from a CSV export.")
    p.add_argument("csv", type=Path, help="Input CSV file")
    p.add_argument("-o", "--output", type=Path, help="Output Excel file (default: reports/report_YYYY-MM-DD.xlsx)")
    p.add_argument("-t", "--title", default="Sales report", help="Report title")
    p.add_argument("--send", action="store_true", help="Email the report (settings in .env)")
    p.add_argument("--dry-run", action="store_true", help="With --send: build the email without sending it")
    p.add_argument("-v", "--verbose", action="store_true")
    args = p.parse_args(argv)

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO, format="%(levelname)s  %(message)s")
    _load_dotenv()

    try:
        df = load_sales(args.csv)
        s = summarize(df)
        output = args.output or Path("reports") / f"report_{date.today():%Y-%m-%d}.xlsx"
        build_excel(df, s, output, args.title)
        log.info("Report created: %s (%d sales, revenue %s)", output, s.sales_count, money(s.revenue))

        if args.send:
            cfg = SMTPConfig.from_env()
            msg = build_email(s, output, cfg, args.title)
            if args.dry_run:
                log.info("[dry-run] Email ready for %s — not sent.", msg["To"])
                print("\n" + msg.get_body(preferencelist=("plain",)).get_content())
            else:
                send_email(msg, cfg)
                log.info("Email sent to %s", msg["To"])
    except (FileNotFoundError, InvalidData, RuntimeError) as e:
        log.error("%s", e)
        return 1
    except smtplib.SMTPException as e:
        log.error("Failed to send the email: %s", e)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
